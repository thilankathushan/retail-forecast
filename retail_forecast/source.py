"""Official UCI acquisition and streaming workbook normalization."""
from collections import Counter
from datetime import datetime, timezone
import csv
import hashlib
import json
from pathlib import Path
import shutil
import urllib.request
import zipfile

PAGE = 'https://archive.ics.uci.edu/dataset/502/online+retail+ii'
URL = 'https://archive.ics.uci.edu/static/public/502/online%2Bretail%2Bii.zip'
COLUMNS = ['Invoice', 'StockCode', 'Description', 'Quantity', 'InvoiceDate', 'Price', 'Customer ID', 'Country']
SHEETS = ['Year 2009-2010', 'Year 2010-2011']


def sha256(path):
    with Path(path).open('rb') as stream:
        return hashlib.file_digest(stream, 'sha256').hexdigest()


def download(raw_dir):
    raw_dir = Path(raw_dir)
    raw_dir.mkdir(parents=True, exist_ok=True)
    archive = raw_dir / 'online_retail_ii.zip'
    manifest_path = raw_dir / 'source.json'
    if archive.exists():
        if not manifest_path.exists():
            raise ValueError('Archive exists without provenance; move it aside before download')
        manifest = json.loads(manifest_path.read_text())
        if sha256(archive) != manifest['archive_sha256']:
            raise ValueError('Archive hash mismatch; move it aside before downloading again')
        return manifest
    temporary = archive.with_suffix('.part')
    try:
        with urllib.request.urlopen(URL, timeout=120) as response, temporary.open('wb') as out:
            shutil.copyfileobj(response, out)
        with zipfile.ZipFile(temporary) as bundle:
            if bundle.namelist() != ['online_retail_II.xlsx']:
                raise ValueError('Unexpected UCI archive contents; inspect source before proceeding')
        temporary.replace(archive)
    finally:
        temporary.unlink(missing_ok=True)
    manifest = {'source_url': PAGE, 'download_url': URL,
                'retrieved_at_utc': datetime.now(timezone.utc).isoformat(),
                'license': 'CC BY 4.0', 'license_url': 'https://creativecommons.org/licenses/by/4.0/',
                'attribution': 'Chen, D. (2012). Online Retail II [Dataset]. UCI Machine Learning Repository. https://doi.org/10.24432/C5CG6D',
                'listed_workbook_size': '43.5 MB', 'archive_bytes': archive.stat().st_size,
                'archive_sha256': sha256(archive)}
    manifest_path.write_text(json.dumps(manifest, indent=2) + '\n')
    return manifest


def import_workbook(raw_dir):
    from openpyxl import load_workbook
    raw_dir = Path(raw_dir)
    manifest = json.loads((raw_dir / 'source.json').read_text())
    archive = raw_dir / 'online_retail_ii.zip'
    if sha256(archive) != manifest['archive_sha256']:
        raise ValueError('Archive hash does not match source.json')
    workbook = raw_dir / 'online_retail_II.xlsx'
    with zipfile.ZipFile(archive) as bundle:
        with bundle.open('online_retail_II.xlsx') as src, workbook.open('wb') as dst:
            shutil.copyfileobj(src, dst)
    output = raw_dir / 'transactions.csv'
    counts, codes = Counter(), Counter()
    earliest = latest = None
    book = load_workbook(workbook, read_only=True, data_only=True)
    try:
        if book.sheetnames != SHEETS:
            raise ValueError(f'Unexpected workbook sheets: {book.sheetnames}')
        with output.open('w', newline='') as stream:
            writer = csv.writer(stream)
            writer.writerow(COLUMNS)
            for sheet in book:
                records = sheet.iter_rows(values_only=True)
                if list(next(records)) != COLUMNS:
                    raise ValueError(f'Unexpected columns in {sheet.title}')
                for record in records:
                    counts[sheet.title] += 1
                    values = list(record)
                    stamp = values[4]
                    if isinstance(stamp, datetime):
                        earliest = min(earliest or stamp, stamp)
                        latest = max(latest or stamp, stamp)
                        values[4] = stamp.isoformat(sep=' ')
                    codes[str(values[1])] += 1
                    writer.writerow(values)
    finally:
        book.close()
    manifest.update({'workbook_bytes': workbook.stat().st_size, 'workbook_sha256': sha256(workbook),
                     'normalized_csv_sha256': sha256(output), 'columns': COLUMNS,
                     'sheet_rows': dict(counts), 'input_rows': sum(counts.values()),
                     'min_timestamp': earliest.isoformat() if earliest else None,
                     'max_timestamp': latest.isoformat() if latest else None})
    (raw_dir / 'source.json').write_text(json.dumps(manifest, indent=2) + '\n')
    (raw_dir / 'code_counts.json').write_text(json.dumps(dict(sorted(codes.items())), indent=2) + '\n')
    return manifest
